"""Editable academic report with actual browser screenshots of sanitized Newman reports."""
import json
from pathlib import Path
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
load = lambda path: json.loads((ROOT/path).read_text(encoding='utf-8'))
data = load('entrega/datos.json')
summary = load('evidence/resumen.json')
base = load('evidence/01-base.json')
broken = load('evidence/02-fallo-controlado.json')
fixed = load('evidence/03-corregida.json')
manual = load('evidence/manual.json')
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.72)
section.footer_distance = Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.08
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 2'].font.size = Pt(11)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(8)
    doc.styles[name].paragraph_format.keep_with_next = True
caption = doc.styles['Caption']
caption.font.name = 'Arial'
caption.font.size = Pt(8)
caption.font.color.rgb = RGBColor(60,60,60)
caption.paragraph_format.space_after = Pt(6)
caption.paragraph_format.keep_with_next = False

def p(text, style=None):
    return doc.add_paragraph(text, style=style)
def labeled(label, text):
    para = p('')
    para.add_run(label+': ').bold = True
    para.add_run(text)
    return para
def link(label, url):
    para = p('')
    para.add_run(label+': ').bold=True
    rel = para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'145D93'); props.append(color)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'17'); props.append(size)
    run.append(props); text=OxmlElement('w:t');text.text=url;run.append(text);hyperlink.append(run);para._p.append(hyperlink)
def picture(filename, caption_text, width=6.65):
    para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.paragraph_format.space_after=Pt(3)
    para.add_run().add_picture(str(ROOT/'evidence/capturas'/filename),width=Inches(width))
    p(caption_text,'Caption')
def table():
    t=doc.add_table(rows=1,cols=6)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.4,.65,.85,.65,1.2,.65]
    names=['Corrida','Casos','Aprobados','Fallidos','Aserciones fallidas','Salida']
    for i,name in enumerate(names):t.rows[0].cells[i].text=name
    rows=[['Base','12','12','0','0','0'],['Defecto local','12','8','4','7','1'],['Corregida','12','12','0','0','0']]
    for values in rows:
        cells=t.add_row().cells
        for i,val in enumerate(values):cells[i].text=val
    for r,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF5' if r==0 else ('F5F7FA' if r%2 else 'FFFFFF'));tcPr.append(shade)
            for para in cell.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after=Pt(0)
                for run in para.runs:run.font.size=Pt(9);run.bold=r==0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

doc.core_properties.title='Validación de una API con Postman y OpenAPI'
doc.core_properties.subject='Informe de evidencia funcional y de contrato'
doc.core_properties.author=data['nombre']
doc.core_properties.comments='Apoyo de IA declarado. Resultados procedentes de ejecuciones locales reales.'

p('Validación de una API con Postman y OpenAPI','Title')
labeled('Estudiante',data['nombre']);labeled('Carné',data['carne'])
if data.get('seccion'):labeled('Sección',data['seccion'])
link('Repositorio',data['repositorio']);link('Video',data['video'])
p('La API local de pedidos cumplió los doce casos funcionales y de contrato evaluados. El cambio incompatible produjo fallos reales y, después de restaurar la fuente, la colección volvió a aprobarse desde datos limpios. Este informe presenta las decisiones de prueba y su evidencia.')
p('Aplicación y riesgos seleccionados','Heading 1')
p('Dos usuarios sintéticos crean y consultan pedidos protegidos. GET pagina y filtra por propietario; POST crea un borrador de CUADERNO a Q25 por unidad; PUT confirma el pedido de forma idempotente. Los riesgos son acceso horizontal, entradas inválidas, cambios no persistidos y confirmaciones duplicadas.')
p('Resultados observados','Heading 1')
dt=datetime.fromisoformat(summary['generatedAt'].replace('Z','+00:00'))-timedelta(hours=6)
p(f'Fecha local {dt.strftime("%d/%m/%Y %H:%M:%S")} en America/Guatemala. Cada corrida ejecutó 12 casos distintos, 14 solicitudes principales y 91 aserciones. Las lecturas auxiliares completan 24 interacciones de pedidos; no se cuentan como casos nuevos.')
table()
p('Reproducción y alcance de la validación','Heading 2')
p('Requisitos verificados: Node 24.19.0, npm 11.17.0 y Newman 6.2.2. Formatos: colección Postman v2.1.0 y OpenAPI 3.1.0. Ejecutar npm ci y npm run verify reproduce base, fallo y corrección. Para usar el Runner: npm run prepare:data, npm start e importar colección y ambiente de postman/; ejecutar las 14 solicitudes en orden, con una iteración y sin persistir variables.')
p('Postman verifica estados, cuerpos, reglas y efectos. Swagger Parser valida el documento OpenAPI y Ajv 2020 valida los cuerpos observados por operación y estado. Esto no constituye validación integral de todas las peticiones, cabeceras y parámetros OpenAPI. No se cubren carga ni autenticación productiva.')
p('Uso de IA y revisión antes de entregar','Heading 2')
p('Codex apoyó diseño, código, contrato y pruebas, y ejecutó los resultados mostrados. El estudiante aportó capturas originales de su ejecución en terminal; también se conserva el visor de informes para el fallo controlado. Los datos de prueba son sintéticos y las sesiones están excluidas. Falta completar la declaración personal de comprensión y revisión.')
p('Los enlaces de repositorio y video son los proporcionados. Antes de exportar el PDF y entregar en Canvas, comprobar que el repositorio contiene los archivos y que el docente puede reproducir el video de máximo tres minutos.')

doc.add_page_break()
p('Evidencia de ejecución y permisos','Heading 1')
picture('terminal/05.png','Figura 1. Captura original aportada por el estudiante: TC09 rechaza lectura y confirmación ajenas; TC10 verifica 404 y TC11 consulta el cambio.',width=6.0)
picture('terminal/07.png','Figura 2. Captura original: 25 solicitudes, 14 scripts y 91 aserciones sin fallos. SETUP, CLEANUP y lecturas auxiliares no son casos nuevos.',width=6.0)

doc.add_page_break()
p('Incompatibilidad detectada y corrección','Heading 1')
p('El defecto real convierte total a texto en GET /orders/{id}, mientras OpenAPI exige un número. TC04 falla aunque HTTP sea 200. También fallan lecturas posteriores de TC09, TC11 y TC12. Se conserva el parche y se restaura la fuente sin modificar el contrato.')
picture('02-fallo-controlado.jpg','Figura 3. Cambio local incompatible y fallo de la aserción de esquema TC04.',width=6.0)
p('Efecto final idempotente','Heading 2')
p('La corrección se ejecutó desde SETUP limpio: 12/12 casos aprobados. En TC12, repetir el PUT conserva fecha, contador persistido 1, Q50 y el mismo conjunto de cuatro pedidos propios. Las lecturas posteriores verifican el efecto final.')
picture('03-corregida.jpg','Figura 4. Resumen de la corrida corregida y extracto de TC04, TC11 y TC12 con lectura final persistida.',width=6.0)

p('Idempotencia, seguridad y pendientes','Heading 1').paragraph_format.page_break_before = True
picture('terminal/06.png','Figura 5. Captura original del estudiante: repetición de PUT, lectura final, verificación del listado y limpieza.',width=6.0)
p('Revisión de dependencias','Heading 2')
p('La instalación original mostró 19 avisos (uno crítico). Después de actualizar dependencias compatibles, npm audit registra 7 entradas: 6 altas y 1 moderada, sin críticas. Pertenecen al árbol de Newman; la API no usa dependencias npm de producción. Esta reducción no equivale a resolver todos los avisos. Se conserva el resultado íntegro en evidence/security/npm-audit.json.')
p('Las capturas de terminal documentan la ejecución anterior a esta actualización. La colección volvió a ejecutarse después del cambio: base y corrección aprobaron 12/12 y 91 aserciones; el defecto produjo 4 casos fallidos. Los informes JSON/TXT contienen las fechas actuales.')
p('Pendiente antes de entregar','Heading 2')
p('Resolver los avisos de seguridad restantes antes de publicar bajo la condición solicitada; ejecutar GitHub Actions y añadir su captura real; comprobar acceso y duración del video; completar revisión personal; exportar este Word a PDF y entregar en Canvas. No se presentan capturas de Actions porque todavía no existe una ejecución publicada verificada.')

p('Colección importada en Postman','Heading 1').paragraph_format.page_break_before = True
p('Las siguientes capturas originales del estudiante muestran la colección importada y sus solicitudes configuradas. Complementan la evidencia de Newman. No muestran una ejecución del Collection Runner, respuestas ni aserciones aprobadas. En ambas aparece No environment; falta seleccionar el ambiente local antes de ejecutar.')
picture('postman/01.png','Figura 6. Postman: colección importada con SETUP, TC01 a TC12 y CLEANUP; instrucciones de uso y alcance de los esquemas.',width=6.5)
picture('postman/07.png','Figura 7. Postman: solicitudes PUT de TC11 y TC12 y DELETE de limpieza configuradas con variables. La configuración por sí sola no demuestra el efecto idempotente.',width=6.5)

footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Página ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
# The supplied default Word template may carry a title paragraph border.
for root in [doc._element,doc.styles.element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
out=ROOT/'entrega/informe.docx'
doc.save(out)
print(out)
